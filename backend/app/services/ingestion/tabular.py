"""CSV / Excel connector for departmental tables."""

from __future__ import annotations

import csv
from io import BytesIO, StringIO
from typing import Any

from openpyxl import load_workbook

from app.core.exceptions import ValidationError
from app.services.ingestion.base import SourceConnector
from app.services.ingestion.payload import ParsedDocument, RawPayload


class TabularConnector(SourceConnector):
    connector_type = "tabular"

    def parse(self, payload: RawPayload) -> ParsedDocument:
        name = payload.final_url.lower()
        mime = payload.mime_type.lower()
        if name.endswith(".xlsx") or name.endswith(".xls") or "spreadsheet" in mime:
            rows = _parse_xlsx(payload.content)
        else:
            rows = _parse_csv(payload.content)
        if not rows:
            raise ValidationError("Tabular source contained no data rows")
        text = "\n".join(
            ", ".join(f"{key}={value}" for key, value in row.items() if value) for row in rows[:50]
        )
        return ParsedDocument(payload=payload, text=text, rows=rows, json_data=rows)


def _parse_csv(content: bytes) -> list[dict[str, Any]]:
    text = content.decode("utf-8-sig", errors="replace")
    reader = csv.DictReader(StringIO(text))
    return [{str(k).strip(): (v or "").strip() for k, v in row.items() if k} for row in reader]


def _parse_xlsx(content: bytes) -> list[dict[str, Any]]:
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise ValidationError(f"Excel workbook could not be read: {type(exc).__name__}") from exc
    sheet = workbook.active
    rows_iter = sheet.iter_rows(values_only=True)
    header_row = next(rows_iter, None)
    if not header_row:
        return []
    headers = [str(cell).strip() if cell is not None else "" for cell in header_row]
    rows: list[dict[str, Any]] = []
    for raw in rows_iter:
        item = {}
        for index, header in enumerate(headers):
            if not header:
                continue
            value = raw[index] if index < len(raw) else None
            item[header] = "" if value is None else str(value).strip()
        if any(item.values()):
            rows.append(item)
    return rows
